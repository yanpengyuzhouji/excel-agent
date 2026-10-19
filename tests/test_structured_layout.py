"""File -> layout -> validation -> SQLite -> catalog regression checks."""
from contextlib import closing
from copy import deepcopy
import sqlite3
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook
from app.service.parser import parse_workbook
from app.service.materializer import build_database
from app.service.service import WorkbookImportService


class StructuredLayoutTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)

    def parse(self, rows, configure=None):
        book = Workbook()
        sheet = book.active
        sheet.title = '统计'
        for row in rows:
            sheet.append(row)
        if configure:
            configure(sheet)
        path = Path(self.temp.name) / 'source.xlsx'
        book.save(path)
        book.close()
        return parse_workbook(str(path))

    def build(self, draft):
        path = str(Path(self.temp.name) / 'result.sqlite')
        build_database(draft, path, import_id='i', workbook_id='w', version=1,
                       file_name='source.xlsx', source_hash='hash')
        return path

    def test_partitioned_report_keeps_metrics_and_separate_tables(self):
        draft = self.parse([
            ['=== 扫描概览 ==='], [], ['扫描时间', '2026-06-24 14:36:22'],
            ['扫描目录', 'D:/business'], ['总文件数', '50,571'], ['损坏率', '58.20%'],
            [], ['=== 损坏类型分布 ==='], [], ['损坏类型', '数量', '严重程度'],
            ['ZIP损坏', 9, '严重'], ['零字节', 8, '严重'],
            [], ['=== 文件类型分布 ==='], [], ['后缀', '数量'], ['.pdf', 100], ['.xlsx', 2]])
        validation = WorkbookImportService._validate_draft(draft)
        self.assertEqual([], validation['errors'])
        self.assertEqual([], [w for w in validation['warnings'] if w['code'].startswith('FORMULA')])
        tables = draft['sheets']
        self.assertEqual([4, 2, 2], [s['row_count'] for s in tables])
        self.assertEqual(['扫描概览', '损坏类型分布', '文件类型分布'], [s['region_title'] for s in tables])
        with closing(sqlite3.connect(self.build(draft))) as db:
            metrics = db.execute(f'SELECT metric, value_raw, numeric_value, unit, _source_row FROM "{tables[0]["table_name"]}"').fetchall()
            self.assertIn(('总文件数', '50,571', 50571.0, None, 5), metrics)
            self.assertIn(('损坏率', '58.20%', .582, '%', 6), metrics)
            self.assertIn(('扫描时间', '2026-06-24 14:36:22', None, None, 3), metrics)

    def test_section_names_disambiguate_identical_metric_labels(self):
        draft = self.parse([['--- 内容重复 ---'], ['涉及文件数', '36,974'], [],
                            ['--- 同名文件 ---'], ['涉及文件数', '130']])
        self.assertEqual([], WorkbookImportService._validate_draft(draft)['errors'])
        self.assertEqual(['内容重复', '同名文件'], [s['region_title'] for s in draft['sheets']])
        self.build(draft)

    def test_multilevel_headers_are_flattened_with_paths(self):
        def configure(sheet):
            for area in ['A1:A2', 'B1:C1', 'D1:E1']:
                sheet.merge_cells(area)
        draft = self.parse([['项目', '计划', None, '实际', None],
                            [None, '数量', '金额', '数量', '金额'],
                            ['甲', 1, 10, 2, 20], ['乙', 3, 30, 4, 40]], configure)
        self.assertEqual([], WorkbookImportService._validate_draft(draft)['errors'])
        sheet = draft['sheets'][0]
        self.assertEqual(['项目', '计划_数量', '计划_金额', '实际_数量', '实际_金额'], [c['display_name'] for c in sheet['columns']])
        self.assertEqual(['B1', 'B2'], sheet['columns'][1]['header_cells'])
        self.build(draft)

    def test_missing_numeric_formula_remains_a_warning(self):
        draft = self.parse([['=== 指标 ==='], ['总数', '=SUM(B5:B6)'], [],
                            ['=== 分布 ==='], ['类型', '数量'], ['A', 1], ['B', 2]])
        self.assertIn('FORMULA_CACHE_MISSING', [w['code'] for w in WorkbookImportService._validate_draft(draft)['warnings']])

    def test_dropped_region_and_tampered_derived_value_are_rejected(self):
        draft = self.parse([['=== 指标 ==='], ['总数', '1,234'], [],
                            ['=== 分布 ==='], ['类型', '数量'], ['A', 1], ['B', 2]])
        broken = deepcopy(draft)
        broken['sheets'].pop()
        with self.assertRaisesRegex(ValueError, 'coverage|覆盖'):
            self.build(broken)
        broken = deepcopy(draft)
        broken['sheets'][0]['rows'][0]['values'][2] = 999
        with self.assertRaisesRegex(ValueError, 'mismatch|派生'):
            self.build(broken)

    def test_unknown_prefix_is_not_silently_classified_as_a_metric(self):
        draft = self.parse([['项目', '金额'], ['甲', 10], [], ['日期', '数量', '金额']]
                            + [['2026-01-01', i, 20] for i in range(6)])
        self.assertTrue(WorkbookImportService._validate_draft(draft)['errors'])

    def test_source_mismatch_is_nonblocking_and_combined_population_is_not_compared(self):
        draft = self.parse([['=== 概览 ==='], ['损坏/可疑文件数', '10'], ['其中可疑文件', '8'], [],
                            ['=== 类型分布 ==='], ['损坏类型','数量','严重程度'], ['A',9,'可疑'], ['B',1,'严重']])
        self.assertEqual([],WorkbookImportService._validate_draft(draft)['warnings'])
        notes=draft['sheets'][0]['quality_notes']
        self.assertEqual(1,len(notes))
        self.assertIn('其中可疑文件',notes[0]['message'])

    def test_three_level_headers_keep_leaf_names(self):
        def configure(sheet):
            for area in ['A1:A3','B1:E1','B2:C2','D2:E2']:
                sheet.merge_cells(area)
        draft=self.parse([['项目','投资',None,None,None],[None,'计划',None,'实际',None],
                          [None,'数量','金额','数量','金额'],['甲',1,10,2,20]],configure)
        self.assertEqual([],WorkbookImportService._validate_draft(draft)['errors'])
        self.assertEqual('投资_实际_金额',draft['sheets'][0]['columns'][4]['display_name'])
        self.build(draft)

    def test_partitioned_table_notes_and_repeated_headers_are_not_records(self):
        draft=self.parse([['=== 分布 ==='],['类型','数量'],['A',1],['B',2],
                          ['类型','数量'],['备注：仅供参考'],['合计',3]])
        sheet=draft['sheets'][0]
        self.assertEqual(2,sheet['row_count'])
        self.assertEqual({'note','summary','repeated_header'},{r['reason'] for r in sheet['excluded_rows']})
        self.assertEqual('INTEGER',sheet['columns'][1]['sql_type'])
        self.build(draft)

    def test_different_schema_inside_one_section_is_not_silently_joined(self):
        draft=self.parse([['=== 分布 ==='],['项目','金额'],['A',1],['B',2],[],
                          ['日期','数量'],['2026-01-01',3],['2026-01-02',4]])
        self.assertTrue(WorkbookImportService._validate_draft(draft)['errors'])

    def test_vertical_merge_cannot_swallow_the_first_numeric_record(self):
        draft=self.parse([['项目','金额'],[None,10],['乙',20]],lambda s:s.merge_cells('A1:A2'))
        self.assertTrue(WorkbookImportService._validate_draft(draft)['errors'])
        self.assertFalse(draft.get('source_layouts'))
        draft=self.parse([['项目','描述'],[None,'首条业务描述'],['乙','第二条']],lambda s:s.merge_cells('A1:A2'))
        self.assertTrue(WorkbookImportService._validate_draft(draft)['errors'])

    def test_uniform_vertical_header_merges_keep_first_business_record(self):
        def configure(sheet):
            sheet.merge_cells('A3:A4')
            sheet.merge_cells('B3:B4')
        draft=self.parse([[None,None],[None,None],['变压器容量(kVA)','接地电缆截面(mm2)'],
                          [None,None],[160,'1*35'],[200,'1*35'],[2500,'1*240']],configure)
        self.assertEqual([],WorkbookImportService._validate_draft(draft)['errors'])
        self.assertEqual([5,6,7],[r['source_row'] for r in draft['sheets'][0]['rows']])
        with closing(sqlite3.connect(self.build(draft))) as db:
            table=draft['sheets'][0]['table_name']
            self.assertEqual(3,db.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0])

    def test_catalog_exposes_region_titles_units_and_quality_notes(self):
        from app.catalog.builder import build_catalog_payload
        draft=self.parse([['=== 概览 ==='],['其中可疑文件','8'],[],['=== 分布 ==='],
                          ['损坏类型','数量','严重程度'],['A',9,'可疑'],['B',1,'严重']])
        payload=build_catalog_payload(self.build(draft),workbook_id='w',display_name='b',source_file='b.xlsx',import_id='i',version=1)
        catalog=payload['payload']
        self.assertIn('概览',catalog['tables'][0]['display_name'])
        self.assertIn('规范数值',[c['display_name'] for c in catalog['tables'][0]['columns']])
        self.assertEqual('SOURCE_SUMMARY_MISMATCH',catalog['tables'][0]['ingestion_notes'][0]['code'])

    def test_parallel_header_regions_preserve_original_column_coordinates(self):
        draft=self.parse([[None,'塔型','设计人',None,'姓名','电话号码'],
                          ['甲组','T1','人员甲',None,'人员乙','tel-a'],
                          [None,'T2','人员丙',None,'人员丁','tel-b']],lambda s:s.merge_cells('A2:A3'))
        self.assertEqual([],WorkbookImportService._validate_draft(draft)['errors'])
        self.assertEqual(2,len(draft['sheets']))
        right=draft['sheets'][1]
        self.assertEqual(['E','F'],right['source_columns'])
        self.assertEqual(['E','F'],[c['source_column'] for c in right['columns']])
        self.assertNotIn('A2',right['cells'])
        self.assertIn('E2',right['cells'])
        self.build(draft)
        broken=deepcopy(draft)
        broken['sheets'][1]['source_columns']=['A','B']
        with self.assertRaisesRegex(ValueError,'mapping'):
            build_database(broken,str(Path(self.temp.name)/'bad.sqlite'),import_id='i',workbook_id='w',
                           version=1,file_name='source.xlsx',source_hash='hash')

    def test_same_row_formula_block_is_not_an_unrelated_parallel_table(self):
        draft=self.parse([['序号','塔型','数量',None,'金额','费用'],
                          [1,'T1',2,None,'=C2*2','=C2*3'],
                          [2,'T2',3,None,'=C3*2','=C3*3']])
        codes=[e['code'] for e in WorkbookImportService._validate_draft(draft)['errors']]
        self.assertNotIn('MULTIPLE_TABLE_REGIONS',codes)
        self.assertEqual(1,len(draft['sheets']))
        from app.service.parser import _formula_bridge
        self.assertFalse(_formula_bridge({'E2':{'raw_value':'=Other!$B2','data_type':'f'}},1,3,5))
        self.assertFalse(_formula_bridge({'E2':{'raw_value':'="B2"','data_type':'f'}},1,3,5))
        self.assertTrue(_formula_bridge({'E2':{'raw_value':'=Other!$B2/$C2','data_type':'f'}},1,3,5))

    def test_explicit_engineering_header_beats_sparse_merged_data(self):
        def configure(s):
            s.merge_cells('A2:A10');s.merge_cells('B2:B10')
        rows=[['序号','塔型','气象区','导线规格','安全系数','水平档距(m)'],
              [1,'T1','A','wire',6.5,80]]+[[None,None,'A','wire',6.5,80] for _ in range(8)]
        draft=self.parse(rows,configure)
        self.assertEqual(1,draft['sheets'][0]['header_row'])
        self.assertEqual([],WorkbookImportService._validate_draft(draft)['errors'])


if __name__ == '__main__':
    unittest.main()
