from contextlib import closing
from decimal import Decimal
import json
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from app.service.materializer import build_database


def draft(value="9007199254740993", sql_type="INTEGER"):
    return {"sheets": [{"name": "Sheet", "enabled": True, "kind": "data",
        "table_name": "business", "columns": [{"enabled": True,
        "physical_name": "amount", "source_column": "A", "sql_type": sql_type}],
        "rows": [{"source_row": 2, "values": [value]}]}]}


class MaterializerIntegrityTests(unittest.TestCase):
    def test_display_rounding_cannot_replace_identifier_value(self):
        data=draft('12','TEXT')
        data['sheets'][0]['cells']={'A2':{'raw_value':12.34,'value':12.34,'display_value':'12'}}
        self.rejected(data)

    def test_legacy_merged_values_reconstruct_blank_source_fact(self):
        data=draft('7','TEXT')
        data['sheets'][0]['merge_lineage']={'A2':{'source_cell':'A1','value':'7'}}
        self.build(data)
        self.assertEqual([('7',)],self.query('SELECT amount FROM business'))

    def test_omitted_source_region_cannot_be_certified(self):
        data=draft('7','TEXT')
        data['sheets'][0]['cells']={'A1':{'raw_value':'old business record','value':'old business record'}}
        data['sheets'][0]['source_classifications']={'2':'data'}
        self.rejected(data)

    def test_merged_formatted_identifier_is_identical_in_database(self):
        from openpyxl import Workbook
        from app.service.parser import parse_workbook
        source=Path(self.storage.name)/'merged.xlsx'
        book=Workbook()
        book.active.append(['项目编号','金额'])
        book.active.append([12,1])
        book.active.append([None,2])
        book.active['A2'].number_format='000000'
        book.active.merge_cells('A2:A3')
        book.save(source)
        book.close()
        data=parse_workbook(str(source))
        self.build(data)
        sheet=data['sheets'][0]
        self.assertEqual([('000012',),('000012',)],self.query(
            f'SELECT "{sheet["columns"][0]["physical_name"]}" FROM "{sheet["table_name"]}" ORDER BY _source_row'))

    def setUp(self):
        self.storage = tempfile.TemporaryDirectory()
        self.addCleanup(self.storage.cleanup)
        self.path = str(Path(self.storage.name) / "book.sqlite")

    def build(self, data):
        build_database(data, self.path, import_id="i", workbook_id="w",
                       version=1, file_name="book.xlsx", source_hash="hash")

    def query(self, sql):
        with closing(sqlite3.connect(self.path)) as db:
            return db.execute(sql).fetchall()

    def rejected(self, data):
        self.path = str(Path(self.storage.name) / f"reject-{len(list(Path(self.storage.name).iterdir()))}.sqlite")
        with self.assertRaises(ValueError):
            self.build(data)
        self.assertEqual([], self.query("SELECT name FROM sqlite_master WHERE type='table'"))

    def test_large_integer_exact(self):
        self.build(draft())
        self.assertEqual([(9007199254740993,)], self.query("SELECT amount FROM business"))

    def test_integer_rejects_loss(self):
        for value in ["1.5", "NaN", "Infinity", "9223372036854775808", "0012", "broken"]:
            with self.subTest(value=value):
                self.rejected(draft(value))

    def test_signed_integer_boundaries(self):
        for value in ["-9223372036854775808", "9223372036854775807", "12.0", "1e3"]:
            with self.subTest(value=value):
                with tempfile.TemporaryDirectory() as folder:
                    self.path = str(Path(folder) / "book.sqlite")
                    self.build(draft(value))
                    self.assertEqual(int(Decimal(value)),
                                     self.query("SELECT amount FROM business")[0][0])

    def test_real_rejects_meaningful_loss(self):
        for value in ["NaN", "Infinity", "1e400", "1e-400", "9007199254740993", "1e23", "0.1234567890123456789"]:
            with self.subTest(value=value):
                self.rejected(draft(value, "REAL"))

    def test_mixed_invalid_does_not_become_null(self):
        data = draft("oops")
        data["sheets"][0]["columns"][0]["mixed_numeric"] = True
        self.rejected(data)

    def test_raw_facts_and_excluded_records(self):
        data = draft("12.30", "REAL")
        sheet = data["sheets"][0]
        sheet.update(source_record_count=2, excluded_rows=[
            {"source_row": 3, "reason": "summary", "values": ["12.30"]}],
            cells={"A1": {"raw_value": "amount", "value": "amount"},
                   "A2": {"raw_value": "=SUM(B2:C2)", "value": "12.30",
                          "data_type": "f", "number_format": "0.00", "display_value": "12.30", "source_cell": "A2"},
                   "A3": {"raw_value": "12.30", "value": "12.30"}})
        self.build(data)
        facts = self.query("SELECT source_cell, detail_json FROM _dbgpt_source_cells ORDER BY source_cell")
        self.assertEqual(["A1", "A2", "A3"], [f[0] for f in facts])
        self.assertEqual("=SUM(B2:C2)", json.loads(facts[1][1])["raw_value"])
        self.assertEqual([("excluded_row",)], self.query("SELECT kind FROM _dbgpt_source_lineage"))
        manifest = json.loads(self.query("SELECT schema_json FROM _dbgpt_import_manifest")[0][0])[0]
        self.assertEqual(1, manifest["reconciliation"]["row_count"])
        self.assertEqual(64, len(manifest["reconciliation"]["sha256"]))

    def test_count_and_source_uniqueness(self):
        data = draft("1")
        data["sheets"][0]["source_record_count"] = 2
        self.rejected(data)
        data = draft("1")
        data["sheets"][0]["rows"] *= 2
        self.rejected(data)

    def test_fact_mismatch_rejected(self):
        data = draft("2")
        data["sheets"][0]["rows"][0]["cells"] = {"A": {"raw_value": "1", "value": "1"}}
        self.rejected(data)

    def test_insert_corruption_rolls_back(self):
        connect = sqlite3.connect

        class CorruptConnection(sqlite3.Connection):
            def executemany(self, sql, parameters):
                if sql.startswith('INSERT INTO "business"'):
                    parameters = [(99, *row[1:]) for row in parameters]
                return super().executemany(sql, parameters)

        with patch("app.service.materializer.sqlite3.connect",
                   side_effect=lambda path: connect(path, factory=CorruptConnection)):
            self.rejected(draft("1"))

    def test_edit_marker_uses_column_id_and_keeps_original(self):
        data = draft("2")
        sheet = data["sheets"][0]
        sheet["columns"][0]["column_id"] = "sheet_1_column_1"
        sheet["rows"][0].update(raw_values=["1"], edited_columns=["sheet_1_column_1"])
        self.build(data)
        self.assertEqual([(2,)], self.query("SELECT amount FROM business"))
        detail = json.loads(self.query("SELECT detail_json FROM _dbgpt_source_cells")[0][0])
        self.assertEqual("1", detail["raw_value"])

    def test_merge_uses_lineage_preserving_blank_original(self):
        data = draft("7")
        sheet = data["sheets"][0]
        sheet["rows"][0]["raw_values"] = [None]
        sheet["merge_lineage"] = {"A2": {"source_cell": "A1", "value": "7"}}
        sheet["cells"] = {"A1": {"raw_value": "7", "value": "7"}}
        self.build(data)
        self.assertEqual([(7,)], self.query("SELECT amount FROM business"))
        detail = json.loads(self.query("SELECT detail_json FROM _dbgpt_source_cells WHERE source_cell='A2'")[0][0])
        self.assertIsNone(detail["raw_value"])

    def test_conflicting_coordinate_facts_rejected(self):
        data = draft("1")
        sheet = data["sheets"][0]
        sheet["cells"] = {"A2": {"value": "1"}}
        sheet["rows"][0]["cells"] = {"A": {"value": "2"}}
        self.rejected(data)

    def test_disabled_and_nonbusiness_facts_preserved(self):
        data = draft("0012", "TEXT")
        sheet = data["sheets"][0]
        sheet["enabled"] = False
        sheet["cells"] = {"A1": {"raw_value": "identifier", "value": "identifier"}}
        self.build(data)
        self.assertEqual(2, self.query("SELECT COUNT(*) FROM _dbgpt_source_cells")[0][0])
        self.assertEqual([], self.query("SELECT name FROM sqlite_master WHERE name='business'"))

    def test_excluded_and_business_source_overlap_rejected(self):
        data = draft("1")
        data["sheets"][0]["excluded_rows"] = [{"source_row": 2, "reason": "note", "values": ["1"]}]
        self.rejected(data)

    def test_text_identifiers_and_numeric_placeholders(self):
        data = draft("0012", "TEXT")
        self.build(data)
        self.assertEqual([("0012",)], self.query("SELECT amount FROM business"))
        self.path = str(Path(self.storage.name) / "mixed.sqlite")
        data = draft("n/a", "integer")
        data["sheets"][0]["columns"][0]["mixed_numeric"] = True
        self.build(data)
        self.assertEqual([(None, "n/a")], self.query("SELECT amount_value, amount_raw FROM business"))

    def test_text_display_normalization_requires_fact_match(self):
        data = draft("000012", "TEXT")
        data["sheets"][0]["cells"] = {"A2": {"raw_value": 12, "value": 12, "display_value": "000012"}}
        self.build(data)
        self.assertEqual([("000012",)], self.query("SELECT amount FROM business"))
        data["sheets"][0]["rows"][0]["values"] = ["000013"]
        self.rejected(data)

    def test_real_money_preserves_original_scale(self):
        data = draft("0.1000", "REAL")
        data["parser"] = {"engine": "openpyxl", "fallback": False}
        data["sheets"][0]["columns"][0].update(semantic_type="money", unit_hint="CNY")
        self.build(data)
        self.assertEqual([(0.1,)], self.query("SELECT amount FROM business"))
        detail = json.loads(self.query("SELECT detail_json FROM _dbgpt_source_cells")[0][0])
        self.assertEqual("0.1000", detail["raw_value"])
        manifest = json.loads(self.query("SELECT schema_json FROM _dbgpt_import_manifest")[0][0])[0]
        self.assertEqual(data["parser"], manifest["parser"])
        self.assertEqual("money", manifest["columns"][0]["semantic_type"])
        self.assertIn("approximate", manifest["columns"][0]["numeric_precision_note"])
        self.assertNotIn("numeric_precision_note", data["sheets"][0]["columns"][0])

    def test_current_parser_xlsx_contract_materializes(self):
        from openpyxl import Workbook
        from app.service.parser import parse_workbook

        source = Path(self.storage.name) / "source.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "data"
        sheet.append(["code", "amount"])
        sheet.append([12, 0.1])
        sheet["A2"].number_format = "000000"
        sheet.append([13, 0.2])
        sheet["A3"].number_format = "000000"
        sheet.append(["合计", 0.3])
        workbook.save(source)
        workbook.close()
        data = parse_workbook(str(source))
        self.build(data)
        table = data["sheets"][0]["table_name"]
        self.assertEqual([("000012", 0.1), ("000013", 0.2)],
                         self.query(f'SELECT code, amount FROM "{table}" ORDER BY _source_row'))
        manifest = json.loads(self.query("SELECT schema_json FROM _dbgpt_import_manifest")[0][0])[0]
        self.assertEqual(3, manifest["source_record_count"])
        self.assertEqual(1, manifest["excluded_row_count"])
        self.assertEqual("openpyxl", manifest["parser"]["engine"])

    def test_sql_coordinate_fact_and_schema_corruption(self):
        connect = sqlite3.connect
        for mode in ["coordinate", "drop_row", "fact", "schema"]:
            with self.subTest(mode=mode):
                class CorruptConnection(sqlite3.Connection):
                    def executemany(self, sql, parameters):
                        parameters = list(parameters)
                        if sql.startswith('INSERT INTO "business"'):
                            if mode == "coordinate":
                                parameters = [(r[0], r[1], 99, r[3]) for r in parameters]
                            elif mode == "drop_row":
                                parameters = []
                        if mode == "fact" and sql.startswith("INSERT INTO _dbgpt_source_cells"):
                            parameters = [(r[0], r[1], r[2], '{}') for r in parameters]
                        return super().executemany(sql, parameters)

                    def execute(self, sql, parameters=()):
                        if mode == "schema" and sql.startswith('CREATE TABLE "business"'):
                            sql = sql.replace('"amount" INTEGER', '"amount" REAL')
                        return super().execute(sql, parameters)

                with patch("app.service.materializer.sqlite3.connect",
                           side_effect=lambda path: connect(path, factory=CorruptConnection)):
                    self.rejected(draft("1"))


if __name__ == "__main__":
    unittest.main()
