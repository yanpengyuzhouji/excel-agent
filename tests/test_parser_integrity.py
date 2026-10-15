"""Golden file-to-draft checks for the automatic Excel ingestion boundary."""
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from app.service.parser import parse_workbook


class ParserIntegrityTests(unittest.TestCase):
    def parse_rows(self, rows, configure=None, overrides=None):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'input.xlsx'
            book = Workbook()
            for row in rows:
                book.active.append(row)
            if configure:
                configure(book.active)
            book.save(path)
            book.close()
            return parse_workbook(str(path),sheet_overrides=overrides)

    def test_text_identifiers_keep_precision_and_leading_zeroes(self):
        draft = self.parse_rows([['项目编号','金额']]+[[str(9007199254740993+i),i+1] for i in range(6)])
        self.assertEqual('TEXT', draft['sheets'][0]['columns'][0]['sql_type'])
        self.assertEqual('9007199254740993', draft['sheets'][0]['rows'][0]['values'][0])
        draft = self.parse_rows([['项目编号','金额']]+[[f'{i:05d}',i+1] for i in range(6)])
        self.assertEqual('TEXT', draft['sheets'][0]['columns'][0]['sql_type'])
        self.assertEqual('00000', draft['sheets'][0]['rows'][0]['values'][0])

    def test_small_table_and_nonbusiness_rows_have_explicit_coverage(self):
        draft = self.parse_rows([['项目','金额'],['甲',10],['乙',20]])
        self.assertEqual('data', draft['sheets'][0]['kind'])
        self.assertEqual(2, draft['sheets'][0]['row_count'])
        draft = self.parse_rows([['项目','金额']]+[[f'P{i}',i+1] for i in range(6)]
            +[['项目','金额'],['合计',21],['备注：测试口径',None]])
        sheet = draft['sheets'][0]
        self.assertEqual(6, sheet['row_count'])
        self.assertEqual({'repeated_header','summary','note'}, {r['reason'] for r in sheet['excluded_rows']})
        self.assertEqual(9, sheet['source_record_count'])

    def test_parallel_tables_are_blocked_instead_of_joined_by_position(self):
        from app.service.service import WorkbookImportService
        draft = self.parse_rows([['项目','金额','项目','金额']]+[[f'P{i}',i,f'Q{i}',i] for i in range(6)])
        self.assertIn('MULTIPLE_TABLE_REGIONS', [i['code'] for i in WorkbookImportService._validate_draft(draft)['errors']])
        draft = self.parse_rows([['项目','金额',None,'日期','数量']]+[[f'P{i}',i,None,'2026-01-01',2] for i in range(6)])
        self.assertIn('MULTIPLE_TABLE_REGIONS',[i['code'] for i in WorkbookImportService._validate_draft(draft)['errors']])

    def test_source_cell_facts_preserve_number_format_and_display_identifier(self):
        def configure(sheet):
            sheet['A2'].number_format = '000000'
            sheet['B2'].number_format = '0.00%'
        draft = self.parse_rows([['项目编号','完成率'],[12,.125],['ABC',.5]], configure)
        sheet = draft['sheets'][0]
        self.assertEqual('000012', sheet['rows'][0]['values'][0])
        self.assertEqual(12, sheet['cells']['A2']['raw_value'])
        self.assertEqual('000000', sheet['cells']['A2']['number_format'])
        self.assertEqual('12.50%', sheet['cells']['B2']['display_value'])

    def test_failed_primary_reader_does_not_silently_downgrade(self):
        from unittest.mock import patch
        with patch('openpyxl.load_workbook', side_effect=ValueError('unsupported package')):
            with self.assertRaisesRegex(ValueError, '可靠|解析'):
                self.parse_rows([['项目','金额'],['甲',10],['乙',20]])

    def test_multilevel_and_separated_tables_require_restructuring(self):
        from app.service.service import WorkbookImportService
        def configure(sheet):
            sheet.merge_cells('A1:A2')
            sheet.merge_cells('B1:C1')
        draft = self.parse_rows([['项目','投资',None],[None,'计划','实际']]+[[f'P{i}',i,i] for i in range(6)],configure)
        codes = [i['code'] for i in WorkbookImportService._validate_draft(draft)['errors']]
        self.assertIn('UNVERIFIED_HEADER_STRUCTURE',codes)
        draft = self.parse_rows([['项目','金额']]+[[f'P{i}',i] for i in range(6)]+[[],['日期','数量']]+[['2026-01-01',2]]*6)
        self.assertIn('MULTIPLE_TABLE_REGIONS',[i['code'] for i in WorkbookImportService._validate_draft(draft)['errors']])

    def test_identifier_display_never_rounds_and_merges_use_anchor_padding(self):
        def rounded(sheet):
            sheet['A2'].number_format='0'
        data = self.parse_rows([['项目编号','金额'],[12.34,1],['B',2]],rounded)['sheets'][0]
        self.assertEqual('12.34',data['rows'][0]['values'][0])
        def merged(sheet):
            sheet['A2'].number_format='000000'
            sheet.merge_cells('A2:A3')
        data = self.parse_rows([['项目编号','金额'],[12,1],[None,2]],merged)['sheets'][0]
        self.assertEqual(['000012','000012'],[r['values'][0] for r in data['rows']])

    def test_later_header_does_not_hide_an_earlier_table(self):
        from app.service.service import WorkbookImportService
        draft = self.parse_rows([['项目','金额'],['甲',10],[],['日期','数量','金额']]
            +[['2026-01-01',i,20] for i in range(6)])
        codes={i['code'] for i in WorkbookImportService._validate_draft(draft)['errors']}
        self.assertTrue(codes & {'UNCLASSIFIED_SOURCE_REGION','MULTIPLE_TABLE_REGIONS'})
        draft=self.parse_rows([['项目名称','金额'],['设计',100],[],['日期','数量','金额','地区']]
            +[['2026-01-01',i,20,'甲'] for i in range(6)])
        codes={i['code'] for i in WorkbookImportService._validate_draft(draft)['errors']}
        self.assertTrue(codes & {'UNCLASSIFIED_SOURCE_REGION','MULTIPLE_TABLE_REGIONS'})

    def test_large_formatted_number_keeps_raw_facts(self):
        def configure(sheet):sheet['B2'].number_format='0.00'
        data=self.parse_rows([['项目','金额'],['A',1e30],['B',2]],configure)['sheets'][0]
        self.assertEqual(1e30,data['cells']['B2']['value'])

    def test_header_override_does_not_bypass_parallel_table_guard(self):
        from app.service.service import WorkbookImportService
        data=self.parse_rows([['项目','金额','项目','金额']]+[[f'P{i}',i,f'Q{i}',i] for i in range(6)],
            overrides={'Sheet':{'header_row':1,'data_start_row':2}})
        self.assertIn('MULTIPLE_TABLE_REGIONS',[i['code'] for i in WorkbookImportService._validate_draft(data)['errors']])

    def test_dates_keep_underlying_ooxml_serial_and_epoch(self):
        from datetime import datetime
        data=self.parse_rows([['日期','金额'],[datetime(2020,1,1),10],[datetime(2020,1,2),20]])
        fact=data['sheets'][0]['cells']['A2']
        self.assertEqual('43831',fact['raw_xml_value'])
        self.assertEqual('n',fact['data_type'])
        self.assertEqual('43831',fact['raw_value'])
        self.assertEqual('2020-01-01T00:00:00',fact['value'])
        self.assertIn('1899-12-30',fact['date_epoch'])
        self.assertIsNone(fact['display_value'])
        def configure(sheet):
            from openpyxl.utils.datetime import CALENDAR_MAC_1904
            sheet.parent.epoch=CALENDAR_MAC_1904
        fact=self.parse_rows([['日期','金额'],[datetime(2020,1,1),10],[datetime(2020,1,2),20]],configure)['sheets'][0]['cells']['A2']
        self.assertEqual('42369',fact['raw_value'])
        self.assertIn('1904-01-01',fact['date_epoch'])
