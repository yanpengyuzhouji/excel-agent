"""Deterministic Excel/CSV parsing used by every ingestion channel.

The parser deliberately does not ask an LLM to decide where data lives.  It keeps
source coordinates and workbook layout facts so a reviewer can correct every
proposal before any database is created.
"""

from __future__ import annotations

import csv
import json
import math
import os
import re
import uuid
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path, PurePosixPath
from typing import Any, Dict, Iterable, List, Optional, Tuple
from xml.etree import ElementTree
from zipfile import ZipFile

from .constants import COMMON_COLUMN_NAMES, HEADER_FIELD_WORDS, PLACEHOLDER_VALUES


@dataclass(frozen=True)
class ParserLimits:
    max_rows: int = 100_000
    max_columns: int = 512
    max_cells: int = 2_000_000
    max_cell_chars: int = 32_767
    max_total_text_chars: int = 64 * 1024 * 1024
    max_uncompressed_bytes: int = 256 * 1024 * 1024
    max_zip_ratio: float = 200.0
    max_zip_entries: int = 10_000
    max_merged_cells: int = 200_000
    max_xml_entry_bytes: int = 64 * 1024 * 1024


def json_default(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def normalize_label(value: Any) -> str:
    text = "" if value is None else str(value)
    return re.sub(r"[\s\r\n_:：]+", "", text).strip().lower()


def safe_identifier(label: Any, fallback: str, used: set[str]) -> str:
    normalized = normalize_label(label)
    candidate = COMMON_COLUMN_NAMES.get(normalized)
    if not candidate:
        candidate = re.sub(r"[^A-Za-z0-9]+", "_", str(label or "")).strip("_").lower()
    if not candidate or candidate[0].isdigit():
        candidate = fallback
    base = candidate[:60]
    candidate = base
    suffix = 2
    while candidate in used:
        candidate = f"{base}_{suffix}"
        suffix += 1
    used.add(candidate)
    return candidate


def _is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _serializable(value: Any) -> Any:
    if _is_empty(value):
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def _infer_column(values: Iterable[Any]) -> Dict[str, Any]:
    present = [v for v in values if not _is_empty(v)]
    placeholders = [v for v in present if str(v).strip().lower() in PLACEHOLDER_VALUES]
    candidates = [v for v in present if v not in placeholders]
    numeric = []
    integral = True
    for value in candidates:
        try:
            number = float(value)
            numeric.append(number)
            if not number.is_integer():
                integral = False
        except (TypeError, ValueError):
            pass
    numeric_ratio = len(numeric) / len(candidates) if candidates else 0.0
    if candidates and numeric_ratio == 1.0:
        sql_type = "INTEGER" if integral else "REAL"
    else:
        sql_type = "TEXT"
    return {
        "sql_type": sql_type,
        "non_empty_count": len(present),
        "numeric_count": len(numeric),
        "placeholder_count": len(placeholders),
        "mixed_numeric": bool(numeric and placeholders),
    }


def _header_score(rows: List[List[Any]], row_index: int) -> float:
    row = rows[row_index]
    non_empty = [v for v in row if not _is_empty(v)]
    if len(non_empty) < 2:
        return -1.0
    unique_ratio = len({normalize_label(v) for v in non_empty}) / len(non_empty)
    string_ratio = sum(isinstance(v, str) for v in non_empty) / len(non_empty)
    following = rows[row_index + 1 : row_index + 11]
    compatible = 0
    width = max(i for i, value in enumerate(row) if not _is_empty(value)) + 1
    for data_row in following:
        filled = sum(not _is_empty(v) for v in data_row[:width])
        if filled >= max(2, min(len(non_empty), width) // 2):
            compatible += 1

    # Real headers are short noun labels; data rows in the same columns are
    # usually longer concrete values (equipment names, descriptions).  Bonus for
    # short labels, penalise rows whose values are long prose.
    avg_len = sum(len(str(v)) for v in non_empty) / len(non_empty)
    if avg_len <= 9:
        label_bonus = 2.5
    elif avg_len <= 12:
        label_bonus = 0.0
    else:
        label_bonus = -6.0

    # Hierarchical directory sheets (merged classification columns) put a run
    # of sparse category-label rows directly under the header.  A header that is
    # followed by such a run is very likely the real one.
    sparse_run = 0
    for data_row in rows[row_index + 1 : row_index + 8]:
        sparse_count = sum(not _is_empty(v) for v in data_row)
        if sparse_count <= 2:
            sparse_run += 1
        else:
            break
    sparse_bonus = min(sparse_run, 4) * 1.5

    # Header cells tend to be common field words (大类/中类/序号/备注...).
    field_hits = sum(1 for v in non_empty if str(v).strip() in HEADER_FIELD_WORDS)
    field_bonus = field_hits * 0.6

    return (
        len(non_empty) * 1.4
        + unique_ratio
        + string_ratio
        + compatible * 1.8
        + label_bonus
        + sparse_bonus
        + field_bonus
    )


def _extract_metadata(rows: List[List[Any]], sheet_name: str) -> List[Dict[str, Any]]:
    result: List[Dict[str, Any]] = []
    seen: set[Tuple[str, str]] = set()
    for row_index, row in enumerate(rows, 1):
        for col_index, value in enumerate(row, 1):
            if _is_empty(value):
                continue
            text = str(value).strip()
            key = None
            metadata_value = None
            if "：" in text or ":" in text:
                parts = re.split(r"[：:]", text, maxsplit=1)
                if len(parts) == 2 and parts[0].strip() and parts[1].strip():
                    key, metadata_value = parts[0].strip(), parts[1].strip()
            if key is None and isinstance(value, str) and col_index < len(row):
                next_value = row[col_index]
                normalized = normalize_label(value)
                if normalized in {
                    "项目名称",
                    "专业类别",
                    "设计阶段",
                    "项目编码",
                    "设计编号",
                    "图号",
                    "审批",
                    "审核",
                    "校核",
                    "设计",
                } and not _is_empty(next_value):
                    key, metadata_value = value.strip(), next_value
            if key is not None:
                marker = (str(key), str(metadata_value))
                if marker not in seen:
                    result.append(
                        {
                            "key": str(key),
                            "physical_name": safe_identifier(
                                key, f"metadata_{len(result) + 1}", set()
                            ),
                            "value": _serializable(metadata_value),
                            "source_sheet": sheet_name,
                            "source_cell": f"{_column_letter(col_index)}{row_index}",
                        }
                    )
                    seen.add(marker)
    return result


def _column_letter(number: int) -> str:
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _rows_from_csv(path: str, limits: ParserLimits) -> List[List[Any]]:
    with open(path, "rb") as binary:
        raw = binary.read(64 * 1024)
    encoding = "utf-8-sig"
    for candidate in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            raw.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    with open(path, "r", encoding=encoding, newline="") as stream:
        rows: List[List[Any]] = []
        cells = 0
        text_chars = 0
        for row_number, row in enumerate(csv.reader(stream), 1):
            if row_number > limits.max_rows:
                raise ValueError(f"CSV exceeds row limit of {limits.max_rows}")
            if len(row) > limits.max_columns:
                raise ValueError(
                    f"CSV row {row_number} exceeds column limit of "
                    f"{limits.max_columns}"
                )
            cells += len(row)
            if cells > limits.max_cells:
                raise ValueError(
                    f"CSV exceeds total cell limit of {limits.max_cells}"
                )
            for value in row:
                if len(value) > limits.max_cell_chars:
                    raise ValueError(
                        f"CSV cell exceeds {limits.max_cell_chars} characters"
                    )
                text_chars += len(value)
            if text_chars > limits.max_total_text_chars:
                raise ValueError(
                    "CSV text exceeds total character limit of "
                    f"{limits.max_total_text_chars}"
                )
            rows.append(list(row))
        return rows


def _parse_cell_reference(reference: str) -> Tuple[int, int]:
    letters = re.match(r"[A-Za-z]+", reference)
    digits = re.search(r"\d+", reference)
    if not letters or not digits:
        raise ValueError(f"Invalid cell reference: {reference}")
    column = 0
    for char in letters.group(0).upper():
        column = column * 26 + ord(char) - 64
    return int(digits.group(0)), column


def _sheet_xml_targets(archive: ZipFile) -> List[Tuple[str, str]]:
    """Resolve each worksheet name to its XML part inside the OOXML package."""
    spreadsheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    office_rel_ns = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    )
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    ns = {"m": spreadsheet_ns, "r": office_rel_ns}
    workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    relationships = ElementTree.fromstring(
        archive.read("xl/_rels/workbook.xml.rels")
    )
    targets = {
        node.attrib["Id"]: node.attrib["Target"]
        for node in relationships.findall(f"{{{package_rel_ns}}}Relationship")
    }
    result: List[Tuple[str, str]] = []
    for sheet_node in workbook.find("m:sheets", ns) or []:
        relationship_id = sheet_node.attrib[f"{{{office_rel_ns}}}id"]
        target = targets[relationship_id]
        if target.startswith("/"):
            target = target.lstrip("/")
        elif not target.startswith("xl/"):
            target = f"xl/{target}"
        result.append((sheet_node.attrib["name"], target))
    return result


def _merge_area(reference: str) -> int:
    start, end = reference.split(":") if ":" in reference else (reference, reference)
    min_row, min_col = _parse_cell_reference(start)
    max_row, max_col = _parse_cell_reference(end)
    return (max_row - min_row + 1) * (max_col - min_col + 1)


def _validated_merge_area(reference: str, limits: ParserLimits) -> int:
    start, end = reference.split(":") if ":" in reference else (reference, reference)
    min_row, min_col = _parse_cell_reference(start)
    max_row, max_col = _parse_cell_reference(end)
    if min_row < 1 or max_row > limits.max_rows:
        raise ValueError(f"Merged range {reference!r} exceeds row limits")
    if min_col < 1 or max_col > limits.max_columns:
        raise ValueError(f"Merged range {reference!r} exceeds column limits")
    return _merge_area(reference)


def _validate_zip_package(path: str, limits: ParserLimits) -> None:
    """Reject oversized/encrypted/pathological OOXML before decompression."""
    with ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > limits.max_zip_entries:
            raise ValueError(
                f"Workbook ZIP has {len(infos)} entries; limit is "
                f"{limits.max_zip_entries}"
            )
        total_uncompressed = 0
        for info in infos:
            path_obj = PurePosixPath(info.filename)
            if path_obj.is_absolute() or ".." in path_obj.parts or "\\" in info.filename:
                raise ValueError(f"Unsafe workbook ZIP entry: {info.filename}")
            if info.flag_bits & 0x1:
                raise ValueError("Encrypted workbook ZIP entries are not supported")
            total_uncompressed += info.file_size
            if total_uncompressed > limits.max_uncompressed_bytes:
                raise ValueError(
                    "Workbook uncompressed size exceeds "
                    f"{limits.max_uncompressed_bytes} bytes"
                )
            if info.filename.endswith(".xml") and info.file_size > limits.max_xml_entry_bytes:
                raise ValueError(
                    f"Workbook XML part {info.filename!r} exceeds "
                    f"{limits.max_xml_entry_bytes} bytes"
                )
            ratio = info.file_size / max(info.compress_size, 1)
            if info.file_size > 1024 * 1024 and ratio > limits.max_zip_ratio:
                raise ValueError(
                    f"Workbook ZIP entry {info.filename!r} has suspicious "
                    f"compression ratio {ratio:.1f}"
                )


def _sheet_metadata_from_xml(path: str, limits: ParserLimits) -> Dict[str, Dict[str, Any]]:
    """Stream the small layout facts read_only worksheets hide.

    openpyxl read_only worksheets drop ``merged_cells``, ``row_dimensions`` and
    ``column_dimensions``.  Those are recovered here by walking each sheet's XML
    once, keeping only merge ranges and hidden row/column flags and discarding
    the (potentially huge) cell data, so memory stays flat.
    """
    spreadsheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    row_tag = f"{{{spreadsheet_ns}}}row"
    col_tag = f"{{{spreadsheet_ns}}}col"
    merge_tag = f"{{{spreadsheet_ns}}}mergeCell"
    result: Dict[str, Dict[str, Any]] = {}
    merged_cells = 0
    with ZipFile(path) as archive:
        for name, target in _sheet_xml_targets(archive):
            hidden_rows: List[int] = []
            hidden_columns: List[str] = []
            merged_ranges: List[str] = []
            try:
                with archive.open(target) as stream:
                    for _, element in ElementTree.iterparse(stream, events=("end",)):
                        tag = element.tag
                        if tag == row_tag:
                            if element.attrib.get("hidden") == "1":
                                hidden_rows.append(int(element.attrib["r"]))
                        elif tag == col_tag:
                            if element.attrib.get("hidden") == "1":
                                start_col = int(element.attrib["min"])
                                end_col = int(element.attrib["max"])
                                if start_col < 1 or end_col > limits.max_columns:
                                    raise ValueError(
                                        "Hidden column range exceeds configured limits"
                                    )
                                for col_index in range(
                                    start_col, end_col + 1,
                                ):
                                    hidden_columns.append(
                                        _column_letter(col_index)
                                    )
                        elif tag == merge_tag:
                            reference = element.attrib["ref"]
                            merged_cells += _validated_merge_area(reference, limits)
                            if merged_cells > limits.max_merged_cells:
                                raise ValueError(
                                    "Workbook merged-cell expansion exceeds "
                                    f"{limits.max_merged_cells} cells"
                                )
                            merged_ranges.append(reference)
                        element.clear()
            except KeyError:
                continue
            result[name] = {
                "merged_ranges": merged_ranges,
                "hidden_rows": hidden_rows,
                "hidden_columns": hidden_columns,
            }
    return result


def _grid_value(rows: List[List[Any]], row: int, col: int) -> Any:
    if row <= 0 or col <= 0:
        return None
    if row > len(rows):
        return None
    current = rows[row - 1]
    if col > len(current):
        return None
    return current[col - 1]


def _parse_streamed(
    formula_book: Any, value_book: Any, metadata: Dict[str, Dict[str, Any]],
    limits: ParserLimits,
) -> List[Tuple]:
    """Read cell values with read_only=True (streaming, low memory).

    Returns the same ``(name, rows, merged, hidden_rows, hidden_columns,
    formulas)`` tuples as the raw-OOXML fallback so ``_build_draft`` is shared.
    """

    parsed_sheets: List[Tuple] = []
    total_cells = 0
    total_text_chars = 0
    for formula_sheet, value_sheet in zip(
        formula_book.worksheets, value_book.worksheets
    ):
        name = formula_sheet.title
        formulas: Dict[str, Dict[str, Any]] = {}
        rows: List[List[Any]] = []
        max_col = 0
        for formula_row, value_row in zip(
            formula_sheet.iter_rows(), value_sheet.iter_rows()
        ):
            if len(rows) >= limits.max_rows:
                raise ValueError(
                    f"Sheet {name!r} exceeds row limit of {limits.max_rows}"
                )
            current: List[Any] = []
            for formula_cell, value_cell in zip(formula_row, value_row):
                if formula_cell.data_type == "f":
                    value = _serializable(value_cell.value)
                    formulas[formula_cell.coordinate] = {
                        "formula": str(formula_cell.value),
                        "cached_value": value,
                    }
                else:
                    value = _serializable(formula_cell.value)
                value_chars = len(value) if isinstance(value, str) else 0
                formula_chars = (
                    len(str(formula_cell.value))
                    if formula_cell.data_type == "f" else 0
                )
                if value_chars > limits.max_cell_chars:
                    raise ValueError(
                        f"Cell {formula_cell.coordinate} exceeds "
                        f"{limits.max_cell_chars} characters"
                    )
                total_text_chars += value_chars + formula_chars
                if total_text_chars > limits.max_total_text_chars:
                    raise ValueError(
                        "Workbook text exceeds total character limit of "
                        f"{limits.max_total_text_chars}"
                    )
                if not _is_empty(formula_cell.value) or not _is_empty(
                    value_cell.value
                ):
                    max_col = max(max_col, formula_cell.column)
                current.append(value)
            if len(current) > limits.max_columns:
                raise ValueError(
                    f"Sheet {name!r} exceeds column limit of {limits.max_columns}"
                )
            total_cells += len(current)
            if total_cells > limits.max_cells:
                raise ValueError(
                    f"Workbook exceeds total cell limit of {limits.max_cells}"
                )
            rows.append(current)
        # Trim trailing empty columns so the grid tracks the true data width,
        # not an inflated sheet dimension (e.g. a stray cell far to the right).
        if max_col and rows:
            rows = [row[:max_col] for row in rows]
        sheet_meta = metadata.get(name, {})
        merged: Dict[str, Dict[str, Any]] = {}
        for reference in sheet_meta.get("merged_ranges", []):
            start, end = (
                reference.split(":") if ":" in reference else (reference, reference)
            )
            min_row, min_col = _parse_cell_reference(start)
            max_merge_row, max_merge_col = _parse_cell_reference(end)
            source_value = _grid_value(rows, min_row, min_col)
            if isinstance(source_value, str):
                total_text_chars += _merge_area(reference) * len(source_value)
                if total_text_chars > limits.max_total_text_chars:
                    raise ValueError(
                        "Expanded merged-cell text exceeds total character "
                        f"limit of {limits.max_total_text_chars}"
                    )
            for row_index in range(min_row, max_merge_row + 1):
                for col_index in range(min_col, max_merge_col + 1):
                    coordinate = f"{_column_letter(col_index)}{row_index}"
                    merged[coordinate] = {
                        "range": reference,
                        "source_cell": start,
                        "value": source_value,
                    }
        parsed_sheets.append(
            (
                name,
                rows,
                merged,
                sheet_meta.get("hidden_rows", []),
                sheet_meta.get("hidden_columns", []),
                formulas,
            )
        )
    return parsed_sheets


def _parse_xlsx_without_openpyxl(path: str, limits: ParserLimits) -> List[Tuple]:
    """Small OOXML fallback so review still works in minimal installations."""

    spreadsheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    office_rel_ns = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    )
    ns = {"m": spreadsheet_ns, "r": office_rel_ns}
    with ZipFile(path) as archive:
        shared: List[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
            shared = [
                "".join(node.text or "" for node in item.iterfind(".//m:t", ns))
                for item in root.findall("m:si", ns)
            ]
        result = []
        total_cells = 0
        total_grid_cells = 0
        total_text_chars = 0
        merged_cells = 0
        for name, target in _sheet_xml_targets(archive):
            sheet_root = ElementTree.fromstring(archive.read(target))
            cells: Dict[Tuple[int, int], Any] = {}
            formulas: Dict[str, Dict[str, Any]] = {}
            hidden_rows: List[int] = []
            hidden_columns: List[str] = []
            max_row = 0
            max_col = 0
            for col in sheet_root.findall(".//m:cols/m:col", ns):
                if col.attrib.get("hidden") == "1":
                    start_col = int(col.attrib["min"])
                    end_col = int(col.attrib["max"])
                    if start_col < 1 or end_col > limits.max_columns:
                        raise ValueError(
                            "Hidden column range exceeds configured limits"
                        )
                    for col_index in range(
                        start_col, end_col + 1
                    ):
                        hidden_columns.append(_column_letter(col_index))
            for row_node in sheet_root.findall(".//m:sheetData/m:row", ns):
                row_number = int(row_node.attrib["r"])
                if row_node.attrib.get("hidden") == "1":
                    hidden_rows.append(row_number)
                for cell in row_node.findall("m:c", ns):
                    coordinate = cell.attrib["r"]
                    row_index, col_index = _parse_cell_reference(coordinate)
                    if row_index > limits.max_rows:
                        raise ValueError(
                            f"Sheet {name!r} exceeds row limit of {limits.max_rows}"
                        )
                    if col_index > limits.max_columns:
                        raise ValueError(
                            f"Sheet {name!r} exceeds column limit of {limits.max_columns}"
                        )
                    value_node = cell.find("m:v", ns)
                    formula_node = cell.find("m:f", ns)
                    cell_type = cell.attrib.get("t")
                    value: Any = value_node.text if value_node is not None else None
                    if cell_type == "s" and value is not None:
                        value = shared[int(value)]
                    elif cell_type == "inlineStr":
                        inline = cell.find("m:is", ns)
                        value = (
                            "".join(
                                node.text or ""
                                for node in inline.iterfind(".//m:t", ns)
                            )
                            if inline is not None
                            else None
                        )
                    elif cell_type == "b" and value is not None:
                        value = value == "1"
                    elif value is not None:
                        try:
                            number = float(value)
                            value = int(number) if number.is_integer() else number
                        except ValueError:
                            pass
                    if not _is_empty(value) or formula_node is not None:
                        value_chars = len(value) if isinstance(value, str) else 0
                        formula_chars = len(formula_node.text or "") if formula_node is not None else 0
                        if value_chars > limits.max_cell_chars:
                            raise ValueError(
                                f"Cell {coordinate} exceeds "
                                f"{limits.max_cell_chars} characters"
                            )
                        total_text_chars += value_chars + formula_chars
                        if total_text_chars > limits.max_total_text_chars:
                            raise ValueError(
                                "Workbook text exceeds total character limit of "
                                f"{limits.max_total_text_chars}"
                            )
                        cells[(row_index, col_index)] = _serializable(value)
                        max_row = max(max_row, row_index)
                        max_col = max(max_col, col_index)
                        total_cells += 1
                        if total_cells > limits.max_cells:
                            raise ValueError(
                                f"Workbook exceeds total cell limit of {limits.max_cells}"
                            )
                    if formula_node is not None:
                        formulas[coordinate] = {
                            "formula": f"={formula_node.text or ''}",
                            "cached_value": _serializable(value),
                        }
            total_grid_cells += max_row * max_col
            if total_grid_cells > limits.max_cells:
                raise ValueError(
                    f"Workbook grid exceeds total cell limit of {limits.max_cells}"
                )
            rows = [
                [
                    cells.get((row_index, col_index))
                    for col_index in range(1, max_col + 1)
                ]
                for row_index in range(1, max_row + 1)
            ]
            merged: Dict[str, Dict[str, Any]] = {}
            for merge in sheet_root.findall(".//m:mergeCells/m:mergeCell", ns):
                reference = merge.attrib["ref"]
                merged_cells += _validated_merge_area(reference, limits)
                if merged_cells > limits.max_merged_cells:
                    raise ValueError(
                        "Workbook merged-cell expansion exceeds "
                        f"{limits.max_merged_cells} cells"
                    )
                start, end = (
                    reference.split(":") if ":" in reference else (reference, reference)
                )
                min_row, min_col = _parse_cell_reference(start)
                max_merge_row, max_merge_col = _parse_cell_reference(end)
                source_value = cells.get((min_row, min_col))
                if isinstance(source_value, str):
                    total_text_chars += _merge_area(reference) * len(source_value)
                    if total_text_chars > limits.max_total_text_chars:
                        raise ValueError(
                            "Expanded merged-cell text exceeds total character "
                            f"limit of {limits.max_total_text_chars}"
                        )
                for row_index in range(min_row, max_merge_row + 1):
                    for col_index in range(min_col, max_merge_col + 1):
                        coordinate = f"{_column_letter(col_index)}{row_index}"
                        merged[coordinate] = {
                            "range": reference,
                            "source_cell": start,
                            "value": source_value,
                        }
            result.append(
                (
                    name,
                    rows,
                    merged,
                    hidden_rows,
                    hidden_columns,
                    formulas,
                )
            )
        return result


def parse_workbook(
    path: str,
    file_name: Optional[str] = None,
    max_rows: Optional[int] = None,
    max_columns: int = 512,
    max_cells: int = 2_000_000,
    max_cell_chars: int = 32_767,
    max_total_text_chars: int = 64 * 1024 * 1024,
    max_uncompressed_bytes: int = 256 * 1024 * 1024,
    max_zip_ratio: float = 200.0,
    max_zip_entries: int = 10_000,
    max_merged_cells: int = 200_000,
    max_xml_entry_bytes: int = 64 * 1024 * 1024,
    sheet_overrides: Optional[Dict[str, Dict[str, int]]] = None,
) -> Dict[str, Any]:
    limits = ParserLimits(
        max_rows=max_rows or 100_000,
        max_columns=max_columns,
        max_cells=max_cells,
        max_cell_chars=max_cell_chars,
        max_total_text_chars=max_total_text_chars,
        max_uncompressed_bytes=max_uncompressed_bytes,
        max_zip_ratio=max_zip_ratio,
        max_zip_entries=max_zip_entries,
        max_merged_cells=max_merged_cells,
        max_xml_entry_bytes=max_xml_entry_bytes,
    )
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        rows = _rows_from_csv(path, limits)
        return _build_draft(
            file_name or Path(path).name,
            [("data", rows, {}, [], [], {})],
            sheet_overrides=sheet_overrides,
        )
    if suffix not in {".xlsx", ".xlsm"}:
        raise ValueError("Only .xlsx, .xlsm and .csv files are supported")
    _validate_zip_package(path, limits)

    try:
        from openpyxl import load_workbook
    except ImportError:
        return _build_draft(
            file_name or Path(path).name,
            _parse_xlsx_without_openpyxl(path, limits),
            sheet_overrides=sheet_overrides,
        )

    # openpyxl 3.1.2 (and some 3.1.x) rejects workbooks whose AutoFilter uses a
    # customFilter whose val is free-form text (e.g. Chinese) in read/write mode:
    #   ValueError: Value must be either numerical or a string containing a wildcard
    # The check lives in worksheet/filters.py and only runs when not read_only.
    # read_only mode skips that metadata entirely.
    #
    # Historical note: this function used to load the workbook twice with
    # read_only=False, materialising every cell of every sheet as a Python
    # object (~570 MB for a 60k-row sheet, two copies held at once), which OOM'd
    # the server on large imports.  read_only=True streams cells instead.  The
    # merged/hidden layout read_only drops is recovered from a lightweight
    # streaming pass over the sheet XML (_sheet_metadata_from_xml).
    try:
        metadata = _sheet_metadata_from_xml(path, limits)
        formula_book = load_workbook(path, data_only=False, read_only=True)
        value_book = load_workbook(path, data_only=True, read_only=True)
    except Exception:
        # Any openpyxl quirk (e.g. certain AutoFilter values, odd packages)
        # falls back to the raw-OOXML parser that understands merged cells,
        # hidden rows and formulas without those checks.
        return _build_draft(
            file_name or Path(path).name,
            _parse_xlsx_without_openpyxl(path, limits),
            sheet_overrides=sheet_overrides,
        )
    try:
        dimension_cells = 0
        for sheet in formula_book.worksheets:
            sheet_rows = int(sheet.max_row or 0)
            sheet_columns = int(sheet.max_column or 0)
            if sheet_rows > limits.max_rows:
                raise ValueError(
                    f"Sheet {sheet.title!r} exceeds row limit of {limits.max_rows}"
                )
            if sheet_columns > limits.max_columns:
                raise ValueError(
                    f"Sheet {sheet.title!r} exceeds column limit of "
                    f"{limits.max_columns}"
                )
            dimension_cells += sheet_rows * sheet_columns
            if dimension_cells > limits.max_cells:
                raise ValueError(
                    f"Workbook dimensions exceed total cell limit of {limits.max_cells}"
                )
        parsed_sheets = _parse_streamed(
            formula_book, value_book, metadata, limits
        )
    finally:
        formula_book.close()
        value_book.close()
    return _build_draft(
        file_name or Path(path).name,
        parsed_sheets,
        sheet_overrides=sheet_overrides,
    )


def _build_draft(
    file_name: str,
    parsed_sheets: List[Tuple],
    sheet_overrides: Optional[Dict[str, Dict[str, int]]] = None,
) -> Dict[str, Any]:
    sheets: List[Dict[str, Any]] = []
    workbook_metadata: List[Dict[str, Any]] = []
    sheet_overrides = sheet_overrides or {}
    for sheet_index, (
        name,
        rows,
        merged,
        hidden_rows,
        hidden_columns,
        formulas,
    ) in enumerate(parsed_sheets):
        sheet_id = f"sheet_{sheet_index + 1}"
        override = sheet_overrides.get(sheet_id) or sheet_overrides.get(name)
        if not rows:
            if override:
                raise ValueError(f"Cannot set a header on empty sheet: {name}")
            sheets.append(
                {
                    "sheet_id": sheet_id,
                    "name": name,
                    "kind": "ignored",
                    "enabled": False,
                    "reason": "empty_sheet",
                    "row_count": 0,
                    "column_count": 0,
                    "columns": [],
                    "rows": [],
                }
            )
            continue
        if override:
            header_row = int(override.get("header_row") or 0)
            data_start_row = int(override.get("data_start_row") or 0)
            if header_row < 1 or header_row > len(rows):
                raise ValueError(
                    f"header_row for {name!r} must be between 1 and {len(rows)}"
                )
            if data_start_row <= header_row or data_start_row > len(rows) + 1:
                raise ValueError(
                    f"data_start_row for {name!r} must be greater than header_row "
                    f"and no greater than {len(rows) + 1}"
                )
            header_index = header_row - 1
            data_start_index = data_start_row - 1
        else:
            candidates = range(min(25, len(rows)))
            header_index = max(candidates, key=lambda index: _header_score(rows, index))
            data_start_index = header_index + 1
        header_score = _header_score(rows, header_index)
        header = rows[header_index]
        last_col = 0
        for row in rows[header_index:]:
            for index, value in enumerate(row, 1):
                if not _is_empty(value):
                    last_col = max(last_col, index)
        data_rows = []
        for source_row, row in enumerate(
            rows[data_start_index:], data_start_index + 1
        ):
            effective = list(row[:last_col]) + [None] * max(0, last_col - len(row))
            for col_index in range(1, last_col + 1):
                coordinate = f"{_column_letter(col_index)}{source_row}"
                if _is_empty(effective[col_index - 1]) and coordinate in merged:
                    effective[col_index - 1] = merged[coordinate]["value"]
            if any(not _is_empty(v) for v in effective):
                data_rows.append(
                    {
                        "source_row": source_row,
                        "hidden": source_row in hidden_rows,
                        "values": [_serializable(v) for v in effective],
                    }
                )
        # Very small key/value layouts frequently look tabular around one row (for
        # example a project cover sheet).  Keep them as metadata by default; the
        # reviewer can still change the proposed kind.
        # A CSV is intrinsically one tabular sheet (named ``data`` above), so
        # allow small files. Excel cover/key-value sheets still need enough
        # records to be proposed as business data.
        is_data = bool(override) or (header_score >= 8 and (
            len(data_rows) >= 5 or (name == "data" and len(data_rows) >= 1)
        ))
        if not is_data:
            metadata = _extract_metadata(rows, name)
            workbook_metadata.extend(metadata)
            sheets.append(
                {
                    "sheet_id": sheet_id,
                    "name": name,
                    "kind": "metadata",
                    "enabled": True,
                    "reason": "layout_is_not_tabular",
                    "row_count": len(rows),
                    "column_count": max((len(row) for row in rows), default=0),
                    "metadata": metadata,
                    "columns": [],
                    "rows": [],
                    "hidden_rows": hidden_rows,
                    "hidden_columns": hidden_columns,
                    "merged_ranges": sorted(
                        {value["range"] for value in merged.values()}
                    ),
                    "formulas": formulas,
                }
            )
            continue

        # Titles immediately above a detected table often carry project metadata
        # (for example "项目名称：...").  Preserve it without turning those rows into
        # records.
        heading_metadata = _extract_metadata(rows[:header_index], name)
        workbook_metadata.extend(heading_metadata)

        used: set[str] = set()
        columns: List[Dict[str, Any]] = []
        for col_index in range(1, last_col + 1):
            original_name = header[col_index - 1] if col_index <= len(header) else None
            empty_header = _is_empty(original_name)
            chinese_alias = str(original_name).strip() if not empty_header else ""
            if empty_header and name == "材料表" and col_index == 11:
                chinese_alias = "选用说明"
            fallback = f"column_{_column_letter(col_index).lower()}"
            physical_name = safe_identifier(chinese_alias, fallback, used)
            values = [row["values"][col_index - 1] for row in data_rows]
            inference = _infer_column(values)
            if "编码" in str(original_name or "").replace("\n", "") or str(
                original_name or ""
            ).strip().lower().endswith("id"):
                inference["sql_type"] = "TEXT"
            columns.append(
                {
                    "column_id": f"{sheet_id}_column_{col_index}",
                    "source_column": _column_letter(col_index),
                    "source_name": None if empty_header else str(original_name).strip(),
                    "display_name": chinese_alias
                    or f"未命名列 {_column_letter(col_index)}",
                    "physical_name": physical_name,
                    "sql_type": inference["sql_type"],
                    "nullable": True,
                    "enabled": True,
                    "empty_header": empty_header,
                    **inference,
                }
            )
        table_name = (
            "material_list"
            if "材料" in name
            else safe_identifier(name, f"table_{sheet_index + 1}", set())
        )
        sheets.append(
            {
                "sheet_id": sheet_id,
                "name": name,
                "kind": "data",
                "enabled": True,
                "table_name": table_name,
                "header_row": header_index + 1,
                "data_start_row": data_start_index + 1,
                "row_count": len(data_rows),
                "column_count": last_col,
                "columns": columns,
                "rows": data_rows,
                "hidden_rows": hidden_rows,
                "hidden_columns": hidden_columns,
                "merged_ranges": sorted({value["range"] for value in merged.values()}),
                "merge_lineage": merged,
                "formulas": formulas,
                "metadata": heading_metadata,
            }
        )
    return {
        "format_version": 1,
        "file_name": file_name,
        "sheets": sheets,
        "metadata": workbook_metadata,
    }


def dump_draft(draft: Dict[str, Any], path: str) -> None:
    target = Path(path)
    pending = target.with_name(f".{target.name}.{uuid.uuid4().hex}.tmp")
    try:
        pending.write_text(
            json.dumps(draft, ensure_ascii=False, indent=2, default=json_default),
            encoding="utf-8",
        )
        os.replace(pending, target)
    finally:
        pending.unlink(missing_ok=True)


def load_draft(path: str) -> Dict[str, Any]:
    return json.loads(Path(path).read_text(encoding="utf-8"))
